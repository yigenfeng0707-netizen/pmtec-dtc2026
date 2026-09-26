# -*- coding: utf-8 -*-
"""UJTL/CJCSI 规范数据清洗入库：仓库根目录 xlsx -> pmtec_system/data/ujtl/ CSV+JSON。

来源（本地留存，不入库）：
- CJCSI_3500.02D_结构化查询表.xlsx  （指令条款的结构化摘要，2026-09-15 签署版）
- UJTL_2026_任务数据集.xlsx         （UJTL 数据库 2026-09-15 快照，1169 现行任务 / 12414 衡量指标）

输出为 UTF-8-SIG CSV（Excel 直开）+ 总览表 JSON（混合版式，无法表格化）。
"""
import csv
import datetime
import json
import os

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CJCSI_XLSX = os.path.join(ROOT, "CJCSI_3500.02D_结构化查询表.xlsx")
UJTL_XLSX = os.path.join(ROOT, "UJTL_2026_任务数据集.xlsx")
OUT_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "ujtl")

# sheet 名 -> 输出文件名（首个非空行为表头，全部数据行原样导出）
SHEET_MAP = {
    "条款索引": "cjcsi_articles.csv",
    "职责矩阵": "cjcsi_responsibilities.csv",
    "系统与接口": "cjcsi_systems.csv",
    "新版变化": "cjcsi_revision_diff.csv",
    "参考文件": "cjcsi_references.csv",
    "术语表": "cjcsi_glossary.csv",
    "任务数据集": "ujtl_tasks.csv",
    "衡量指标": "ujtl_measures.csv",
    "变更清单": "ujtl_changes.csv",
    "删除审计": "ujtl_deletions.csv",
    "编号段索引": "ujtl_segments.csv",
    "字段说明": "ujtl_field_spec.csv",
}
OVERVIEW_MAP = {  # 混合版式表 -> JSON
    "总览": "cjcsi_overview.json",
    "总览与统计": "ujtl_overview.json",
}


def cell(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def export_sheet(ws, path):
    rows = [[cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows)
    return len(rows) - 1  # 数据行数（去掉表头）


def export_overview(ws, path, source):
    rows = [[cell(c) for c in r] for r in ws.iter_rows(values_only=True)]
    rows = [[c for c in r if c] for r in rows if any(r)]
    data = {"source": source, "sheet": ws.title, "rows": rows}
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    return len(rows)


def main():
    out = os.path.normpath(OUT_DIR)
    os.makedirs(out, exist_ok=True)
    for xlsx in (CJCSI_XLSX, UJTL_XLSX):
        if not os.path.exists(xlsx):
            raise SystemExit(f"[缺失] 找不到源文件: {xlsx}")
    total = 0
    for src in (CJCSI_XLSX, UJTL_XLSX):
        wb = openpyxl.load_workbook(src, read_only=True)
        for ws in wb.worksheets:
            if ws.title in SHEET_MAP:
                name = SHEET_MAP[ws.title]
                n = export_sheet(ws, os.path.join(out, name))
                print(f"[OK] {name:<28} 数据行 {n}")
                total += n
            elif ws.title in OVERVIEW_MAP:
                name = OVERVIEW_MAP[ws.title]
                n = export_overview(ws, os.path.join(out, name), os.path.basename(src))
                print(f"[OK] {name:<28} 版式行 {n}")
        wb.close()
    print(f"     合计 {total} 数据行 -> {out}")


if __name__ == "__main__":
    main()
